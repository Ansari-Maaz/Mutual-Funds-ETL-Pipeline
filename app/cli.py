from __future__ import annotations
import argparse, json, re
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
import yaml
from app.catalog import country_rows, discovered_rows
from app.connectors.registry import all_connectors
from app.config import timeout_seconds, user_agent
import httpx

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"

def catalogue() -> list[dict]:
    candidates = (yaml.safe_load((ROOT / "config/sources.yaml").read_text()) or []) + discovered_rows()
    verification_path=REPORTS/"source_verification.json"
    if verification_path.exists():
        verified={r["source_id"]:r for r in json.loads(verification_path.read_text())}
        for source in candidates:
            result=verified.get(source["source_id"])
            if result:
                source["verification_status"]=result["verification_status"]
                source["last_checked_at"]=result["checked_at"]
                source["http_status"]=result["http_status"]
                source["usable_fund_data_confirmed"]=result["usable_fund_data_confirmed"]
                if result.get("error"): source["failure_reason"]=result["error"]
    run_path=REPORTS/"latest_run.json"
    if run_path.exists():
        try: run_document=json.loads(run_path.read_text()); run_rows=run_document.get("results",[])
        except (OSError,ValueError): run_document={}; run_rows=[]
    else: run_document={}; run_rows=[]
    run_by_id={row.get("source_id"):row for row in run_rows}
    for source in candidates:
        result=run_by_id.get(source.get("source_id"))
        if not result: continue
        source["records_extracted"]=result.get("records_extracted",0)
        source["records_loaded"]=result.get("records_loaded",0)
        source["nav_observations_loaded"]=result.get("navs_loaded",0)
        source["rejected_records"]=result.get("rejected_records",0)
        if result.get("status")=="Succeeded":
            source["integration_status"]="Integrated"
            source["verification_status"]="Usable Fund Data Confirmed"
            source["usable_fund_data_confirmed"]=True
            source["last_successful_ingestion_at"]=run_document.get("generated_at","")
            source["last_checked_at"]=run_document.get("generated_at","")
            source["failure_reason"]=""
        elif result.get("status")=="Failed":
            source["integration_status"]="Failed"
            source["failure_reason"]=result.get("error", "ETL run failed")
    # Keep the inventory's integration state tied to persisted database rows,
    # even when latest_run.json describes only a single-source execution.
    try:
        from sqlalchemy import func, select
        from app.db.session import SessionLocal
        from app.models.schema import Fund, NavObservation
        with SessionLocal() as session:
            fund_counts=dict(session.execute(select(Fund.source_id,func.count()).group_by(Fund.source_id)).all())
            nav_counts=dict(session.execute(select(Fund.source_id,func.count()).select_from(NavObservation).join(Fund,NavObservation.fund_id==Fund.id).group_by(Fund.source_id)).all())
        for source in candidates:
            source_id=source.get("source_id")
            if fund_counts.get(source_id,0):
                source["integration_status"]="Integrated"
                source["verification_status"]="Usable Fund Data Confirmed"
                source["usable_fund_data_confirmed"]=True
                source["records_loaded"]=fund_counts[source_id]
                source["nav_observations_loaded"]=nav_counts.get(source_id,0)
    except Exception:
        pass
    researched_codes={r["country_code"] for r in candidates}
    return country_rows(researched_codes)+candidates

def verify_sources() -> None:
    result=[]
    for src in catalogue():
        url=src.get("endpoint_url", "")
        row={"source_id":src["source_id"],"country_code":src["country_code"],"url":url,
             "checked_at":datetime.now(timezone.utc).isoformat(),"http_status":None,
             "content_type":None,"response_size":None,"preview":None,
             "verification_status":"Unavailable","usable_fund_data_confirmed":False,"error":None}
        try:
            if urlsplit(url).scheme not in {"http","https"} or not urlsplit(url).netloc:
                raise ValueError("Invalid HTTP(S) endpoint")
            response=httpx.get(url, timeout=timeout_seconds(), follow_redirects=True,
                headers={"User-Agent":user_agent(),"Accept":"application/json, text/csv, application/xml, text/html"})
            body=response.content
            preview=re.sub(r"(?i)(api[_-]?key|token|password)=([^&\s]+)",r"\1=[REDACTED]",body[:500].decode("utf-8",errors="replace"))
            ct=response.headers.get("content-type","")
            row.update(http_status=response.status_code,content_type=ct,response_size=len(body),preview=preview)
            if response.status_code in (401,403): row["verification_status"]="Authentication Required"
            elif response.status_code==429: row["verification_status"]="Blocked"
            elif response.status_code>=400: row["verification_status"]="Unavailable"
            elif "text/html" in ct: row["verification_status"]="Reachable" if "<html" in preview.lower() else "Accessible — No Fund Data"
            else: row["verification_status"]="Reachable"
            # Never infer usable fund data from an HTTP success or homepage response.
        except (httpx.HTTPError, ValueError) as exc:
            row["error"]=f"{type(exc).__name__}: {str(exc)[:300]}"
        result.append(row)
    REPORTS.mkdir(exist_ok=True)
    (REPORTS/"source_verification.json").write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n")
    lines=["# Source verification", "", f"Generated at {datetime.now(timezone.utc).isoformat()}", "",
      "A reachable homepage is not confirmation of fund-level data. No source is marked usable without a provider-specific parser and evidence.", "",
      "| Source | Country | Status | HTTP | Type | Bytes | Error |", "|---|---|---|---:|---|---:|---|"]
    for x in result: lines.append(f"| {x['source_id']} | {x['country_code']} | {x['verification_status']} | {x['http_status'] or ''} | {x['content_type'] or ''} | {x['response_size'] if x['response_size'] is not None else ''} | {(x['error'] or '').replace('|','/')} |")
    (REPORTS/"source_verification.md").write_text("\n".join(lines)+"\n")
    print(f"Checked {len(result)} candidate endpoints; report: reports/source_verification.json")

def generate_workbook() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.datavalidation import DataValidation
    rows=catalogue(); REPORTS.mkdir(exist_ok=True); path=REPORTS/"global_mutual_fund_source_inventory.xlsx"
    wb=Workbook(); dash=wb.active; dash.title="Dashboard"
    dash.append(["Global Mutual Fund ETL — Inventory Summary"])
    dash.append(["Metric","Count"])
    researched_codes={r.get("country_code") for r in rows if r.get("verification_status")!="Undiscovered"}
    dash.append(["Countries in required scope",24]); dash.append(["Countries with candidate research documented",len(researched_codes)]); dash.append(["Countries with discovery pending",24-len(researched_codes)])
    dash.append(["Candidate source records",len(rows)]); dash.append(["Usable sources confirmed",sum(r.get("usable_fund_data_confirmed") is True for r in rows)])
    dash.append(["Integrated sources",sum(r.get("integration_status")=="Integrated" for r in rows)])
    dash.append(["Verification status", "Count"])
    for status in ["Undiscovered","Discovered","Reachable","Authentication Required","Blocked","Unavailable","Accessible — No Fund Data","Usable Fund Data Confirmed"]:
        dash.append([status,sum(r.get("verification_status")==status for r in rows)])
    regional={"Asia-Pacific":("Asia-Pacific",["JP","CN","KR","AU","SG","HK","TW"]),"Europe":("Europe",["GB","DE","FR","IT","CH","NL","DK","SE","LU","IE"]),"North America":("North America",["US","CA","MX"]),"Other Markets":("Other markets",["RU","BR","AE","ZA"])}
    columns=["Source ID","Country","ISO country code","Region","Organization","Source name","Official website","API endpoint or download URL","API documentation URL","Source category","Access method","Authentication required","Data format","Fund master data available","Current NAV available","Historical NAV available","Holdings available","Currency information available","Available identifiers","Historical data coverage","Update frequency","Rate limits","Licensing status","Commercial reuse status","Endpoint verification status","Usable fund data confirmed","ETL implementation status","Integration test status","Last checked timestamp","Last successful ingestion timestamp","Records extracted","Records loaded","Error or blocker","Evidence URL","Research notes"]
    mapping={"Source ID":"source_id","Country":"country_name","ISO country code":"country_code","Region":"region","Organization":"organization","Source name":"source_name","Official website":"official_website","API endpoint or download URL":"endpoint_url","API documentation URL":"documentation_url","Source category":"source_category","Access method":"access_method","Authentication required":"authentication_required","Data format":"response_format","Fund master data available":"fund_master_available","Current NAV available":"current_nav_available","Historical NAV available":"historical_nav_available","Holdings available":"holdings_available","Currency information available":"available_currency_fields","Available identifiers":"available_identifiers","Historical data coverage":"historical_coverage","Update frequency":"update_frequency","Rate limits":"rate_limit_notes","Licensing status":"license_status","Commercial reuse status":"commercial_reuse_status","Endpoint verification status":"verification_status","Usable fund data confirmed":"usable_fund_data_confirmed","ETL implementation status":"integration_status","Integration test status":"integration_test_status","Last checked timestamp":"last_checked_at","Last successful ingestion timestamp":"last_successful_ingestion_at","Records extracted":"records_extracted","Records loaded":"records_loaded","Error or blocker":"failure_reason","Evidence URL":"evidence_url","Research notes":"notes"}
    sheets={"Asia-Pacific":"Asia-Pacific","Europe":"Europe","North America":"North America","Other Markets":"Other markets"}
    for title,region in sheets.items():
        ws=wb.create_sheet(title); ws.append(columns)
        for row in rows:
            if row.get("region")==region: ws.append([row.get(mapping[c], "") for c in columns])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for idx in (7,8,9,34):
                cell=row[idx-1]
                if isinstance(cell.value,str) and cell.value.startswith(("https://","http://")):
                    cell.hyperlink=cell.value; cell.style="Hyperlink"
        if ws.max_row>1:
            table=Table(displayName=title.replace(" ","")+"Sources",ref=ws.dimensions); table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False); ws.add_table(table)
            for column, formula in [(23,'"Permitted,Restricted,Permission Required,Unknown"'),(25,'"Undiscovered,Discovered,Reachable,Authentication Required,Blocked,Unavailable,Accessible — No Fund Data,Usable Fund Data Confirmed"'),(27,'"Not Started,In Progress,Implemented,Tested,Integrated,Failed"')]:
                dv=DataValidation(type="list",formula1=formula,allow_blank=True); ws.add_data_validation(dv); dv.add(f"{ws.cell(2,column).coordinate}:{ws.cell(ws.max_row,column).coordinate}")
            last=ws.max_row
            ws.conditional_formatting.add(f"Y2:Y{last}",FormulaRule(formula=['OR(Y2="Usable Fund Data Confirmed",Y2="Reachable")'],fill=PatternFill("solid",fgColor="C6EFCE")))
            ws.conditional_formatting.add(f"Y2:Y{last}",FormulaRule(formula=['OR(Y2="Blocked",Y2="Unavailable")'],fill=PatternFill("solid",fgColor="FFC7CE")))
            ws.conditional_formatting.add(f"AA2:AA{last}",FormulaRule(formula=['AA2="Integrated"'],fill=PatternFill("solid",fgColor="C6EFCE")))
            ws.conditional_formatting.add(f"AA2:AA{last}",FormulaRule(formula=['OR(AA2="Failed",AA2="Not Started")'],fill=PatternFill("solid",fgColor="FFEB9C")))
        for col in ws.columns:
            letter=col[0].column_letter; ws.column_dimensions[letter].width=min(42,max(14,max(len(str(c.value or "")) for c in col)+2))
        for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D"); cell.alignment=Alignment(wrap_text=True)
    log=wb.create_sheet("Source Verification Log"); log.append(["Source ID","Country","Checked at","HTTP status","Format","Fund data confirmed","Status","Error","Evidence"])
    vp=REPORTS/"source_verification.json"
    if vp.exists():
        for x in json.loads(vp.read_text()): log.append([x.get("source_id"),x.get("country_code"),x.get("checked_at"),x.get("http_status"),x.get("content_type"),x.get("usable_fund_data_confirmed"),x.get("verification_status"),x.get("error"),x.get("url")])
    etl=wb.create_sheet("ETL Integration Status"); etl.append(["Source ID","Country","Connector","Integration status","Last successful run","Records extracted","Records loaded","Errors"])
    registered=all_connectors()
    for r in rows:
        etl.append([r.get("source_id"),r.get("country_code"),
          "Implemented" if r.get("source_id") in registered else "Not implemented",
          r.get("integration_status"),r.get("last_successful_ingestion_at",""),
          r.get("records_extracted",0),r.get("records_loaded",0),r.get("failure_reason","")])
    dd=wb.create_sheet("Data Dictionary"); dd.append(["Field","Meaning / accepted values"])
    for k,v in [("Country code","ISO 3166-1 alpha-2"),("Verification status","Undiscovered; Discovered; Reachable; Authentication Required; Blocked; Unavailable; Accessible — No Fund Data; Usable Fund Data Confirmed"),("Integration status","Not Started; In Progress; Implemented; Tested; Integrated; Failed"),("Licensing status","Permitted; Restricted; Permission Required; Unknown"),("Usable fund data confirmed","True only after actual fund-level records are inspected")]: dd.append([k,v])
    for ws in wb.worksheets:
        ws.freeze_panes="A2"
        for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17365D"); cell.alignment=Alignment(wrap_text=True)
        ws.auto_filter.ref=ws.dimensions
    chart=BarChart(); chart.type="bar"; chart.style=10; chart.title="Sources by verification status"; chart.y_axis.title="Status"; chart.x_axis.title="Sources"
    chart.add_data(Reference(dash,min_col=2,min_row=9,max_row=17),titles_from_data=True); chart.set_categories(Reference(dash,min_col=1,min_row=10,max_row=17)); chart.height=7; chart.width=14; dash.add_chart(chart,"D2")
    wb.save(path); print(path)

def export_catalogue() -> None:
    REPORTS.mkdir(exist_ok=True)
    (REPORTS/"source_catalogue.json").write_text(json.dumps(catalogue(),indent=2,ensure_ascii=False)+"\n")
    if not (REPORTS/"sample_funds.json").exists(): (REPORTS/"sample_funds.json").write_text("[]\n")
    if not (REPORTS/"sample_nav.json").exists(): (REPORTS/"sample_nav.json").write_text("[]\n")

def export_database_samples() -> dict[str,int]:
    from sqlalchemy import select, func
    from app.db.session import SessionLocal
    from app.models.schema import Fund,NavObservation
    REPORTS.mkdir(exist_ok=True)
    with SessionLocal() as session:
        source_ids=session.scalars(select(Fund.source_id).distinct().order_by(Fund.source_id)).all()
        funds_json=[]; nav_json=[]; counts={"funds":0,"nav_observations":0,"by_source":{}}
        for source_id in source_ids:
            funds=session.scalars(select(Fund).where(Fund.source_id==source_id).order_by(Fund.source_fund_id).limit(100)).all()
            nav_rows=session.execute(select(Fund.source_fund_id,Fund.name,NavObservation.valuation_date,
              NavObservation.nav,NavObservation.currency).join(NavObservation,NavObservation.fund_id==Fund.id)
              .where(Fund.source_id==source_id).order_by(NavObservation.valuation_date.desc(),Fund.source_fund_id).limit(1000)).all()
            funds_json.extend({"source_id":f.source_id,"source_fund_id":f.source_fund_id,"name":f.name,
              "domicile_country":f.domicile_country,"currency":f.currency,"isin":f.isin} for f in funds)
            nav_json.extend({"source_id":source_id,"source_fund_id":row[0],"fund_name":row[1],
              "valuation_date":row[2].isoformat(),"nav":str(row[3]),"currency":row[4]} for row in nav_rows)
            fund_count=session.scalar(select(func.count()).select_from(Fund).where(Fund.source_id==source_id)) or 0
            nav_count=session.scalar(select(func.count()).select_from(NavObservation).join(Fund,NavObservation.fund_id==Fund.id).where(Fund.source_id==source_id)) or 0
            counts["by_source"][source_id]={"funds":fund_count,"nav_observations":nav_count}
            counts["funds"]+=fund_count; counts["nav_observations"]+=nav_count
    (REPORTS/"sample_funds.json").write_text(json.dumps(funds_json,indent=2,ensure_ascii=False)+"\n")
    (REPORTS/"sample_nav.json").write_text(json.dumps(nav_json,indent=2,ensure_ascii=False)+"\n")
    return counts

def write_latest_run_report(run: dict,counts: dict[str,int]) -> None:
    from app.config import database_url
    timestamp=run.get("generated_at","")
    rows=run.get("results",[])
    lines=["# Latest execution report","",f"Generated at: {timestamp}","",
      "## Ingestion results","","| Source | Status | Funds extracted | Funds loaded | NAV observations upserted | Rejected rows | Error |",
      "|---|---|---:|---:|---:|---:|---|"]
    for row in rows:
        lines.append(f"| {row.get('source_id','')} | {row.get('status','')} | {row.get('records_extracted',0)} | {row.get('records_loaded',0)} | {row.get('navs_loaded',0)} | {row.get('rejected_records',0)} | {(row.get('error') or '').replace('|','/')} |")
    db_url=database_url()
    storage_note=("Temporary PostgreSQL validation cluster under /tmp; records are not in the configured Docker database." if "/tmp/mfetl-sock" in db_url else "Configured project PostgreSQL database.")
    lines += ["","## Database counts","","| Source | Funds | NAV observations |","|---|---:|---:|"]
    for source_id,source_counts in counts.get("by_source",{}).items():
        lines.append(f"| {source_id} | {source_counts['funds']} | {source_counts['nav_observations']} |")
    lines += [f"| **Total** | **{counts['funds']}** | **{counts['nav_observations']}** |",
      f"- Database used: {storage_note}",
      "- Raw snapshots are ignored by Git; source manifests record URLs, checksums, and sizes.",
      "- Source rejects are recorded in source-specific JSONL files under `reports/`.",
      "- Sample exports contain actual database rows: `sample_funds.json` and `sample_nav.json`.",
      "- Commercial reuse rights require separate review for all providers.",""]
    (REPORTS/"latest_run.md").write_text("\n".join(lines))

def sync_sources() -> None:
    from sqlalchemy.dialects.postgresql import insert
    from app.db.session import SessionLocal
    from app.models.schema import Source
    values=[]
    rows=catalogue()
    researched_codes={r["country_code"] for r in rows if r.get("verification_status")!="Undiscovered"}
    obsolete_ids=[r["source_id"] for r in country_rows() if r["country_code"] in researched_codes]
    for row in rows:
        values.append({"source_id":row["source_id"],"country_code":row["country_code"],
          "organization":str(row.get("organization","Unknown")),"source_name":row.get("source_name","Unknown"),
          "endpoint_url":row.get("endpoint_url") or None,
          "verification_status":row.get("verification_status","Undiscovered"),
          "integration_status":row.get("integration_status","Not Started"),
          "license_status":row.get("license_status","Unknown")})
    with SessionLocal() as session:
        if obsolete_ids:
            from sqlalchemy import delete
            session.execute(delete(Source).where(Source.source_id.in_(obsolete_ids)))
        for row in values:
            stmt=insert(Source).values(**row).on_conflict_do_update(index_elements=[Source.source_id],set_={k:v for k,v in row.items() if k!="source_id"})
            session.execute(stmt)
        session.commit()
    print(f"Synchronized {len(values)} catalog source/country rows.")

def main() -> None:
    parser=argparse.ArgumentParser(prog="mfetl"); sub=parser.add_subparsers(dest="cmd",required=True)
    sub.add_parser("verify-sources"); sub.add_parser("inventory"); sub.add_parser("sync-sources")
    download=sub.add_parser("download"); download.add_argument("--source",default="cvm_open_fund_data"); download.add_argument("--start-date"); download.add_argument("--end-date"); download.add_argument("--refresh",action="store_true")
    run=sub.add_parser("run"); run.add_argument("--all",action="store_true"); run.add_argument("--country"); run.add_argument("--source"); run.add_argument("--start-date"); run.add_argument("--end-date")
    args=parser.parse_args()
    if args.cmd=="verify-sources": verify_sources()
    elif args.cmd=="inventory":
        export_catalogue(); generate_workbook()
    elif args.cmd=="sync-sources": sync_sources()
    elif args.cmd=="download":
        if args.source=="dk_bde_nonmonetary_funds":
            from app.connectors.denmark_bde_funds import download_denmark_fund_list
            print(json.dumps(download_denmark_fund_list(refresh=args.refresh),indent=2))
            return
        if args.source in {"de_bde_nonmonetary_funds","it_bde_nonmonetary_funds","nl_bde_nonmonetary_funds","lu_bde_nonmonetary_funds"}:
            from app.connectors.bde_european_funds import download_bde_fund_list
            country_code=args.source[:2].upper()
            print(json.dumps(download_bde_fund_list(country_code,refresh=args.refresh),indent=2))
            return
        if args.source=="hk_sfc_authorised_unit_trusts":
            from app.connectors.hk_sfc_funds import download_hk_sfc_fund_list
            print(json.dumps(download_hk_sfc_fund_list(refresh=args.refresh),indent=2))
            return
        if args.source=="amf_geco":
            from app.connectors.amf_geco import download_amf_geco
            print(json.dumps(download_amf_geco(refresh=args.refresh),indent=2))
            return
        if args.source=="cbi_authorised_ucits":
            from app.connectors.central_bank_ireland import download_ireland_ucits_register
            print(json.dumps(download_ireland_ucits_register(refresh=args.refresh),indent=2))
            return
        if args.source=="fi_sweden_ucits_holdings":
            from app.connectors.fi_sweden import download_fi_sweden
            print(json.dumps(download_fi_sweden(refresh=args.refresh),indent=2))
            return
        if args.source=="cnmv_spanish_fund_register":
            from app.connectors.cnmv import download_cnmv_fund_register
            print(json.dumps(download_cnmv_fund_register(refresh=args.refresh),indent=2))
            return
        if args.source=="sec_open_end_fund_classes":
            from app.connectors.sec import download_sec_series_classes
            print(json.dumps(download_sec_series_classes(refresh=args.refresh),indent=2))
            return
        if args.source=="fca_authorised_funds":
            from app.connectors.fca import download_fca_fund_register
            print(json.dumps(download_fca_fund_register(refresh=args.refresh),indent=2))
            return
        if args.source=="finma_swiss_fund_list":
            from app.connectors.finma import download_finma_fund_list
            print(json.dumps(download_finma_fund_list(refresh=args.refresh),indent=2))
            return
        if args.source=="asisa_cis_fund_list":
            from app.connectors.asisa import download_asisa_fund_list
            print(json.dumps(download_asisa_fund_list(refresh=args.refresh),indent=2))
            return
        if args.source=="gsfm_unit_prices":
            from app.connectors.gsfm import download_gsfm_prices
            print(json.dumps(download_gsfm_prices(refresh=args.refresh),indent=2))
            return
        if args.source=="csrc_public_fund_index":
            from app.connectors.csrc import download_csrc_index
            print(json.dumps(download_csrc_index(refresh=args.refresh),indent=2))
            return
        if args.source!="cvm_open_fund_data": raise SystemExit(f"No downloader registered for {args.source!r}.")
        from app.connectors.cvm import download_cvm_data
        today=datetime.now(timezone.utc).date()
        start=date.fromisoformat(args.start_date) if args.start_date else date(today.year,today.month,1)
        end=date.fromisoformat(args.end_date) if args.end_date else today
        if start>end: raise SystemExit("--start-date must be on or before --end-date")
        months=[]; year,month=start.year,start.month
        while (year,month)<=(end.year,end.month):
            months.append(f"{year:04d}{month:02d}"); month+=1
            if month==13: year,month=year+1,1
        results=download_cvm_data(months,refresh=args.refresh)
        print(json.dumps(results,indent=2))
    else:
        candidates=all_connectors()
        if args.source:
            candidates={k:v for k,v in candidates.items() if k==args.source}
            if not candidates: raise SystemExit(f"No connector registered for source {args.source!r}; no ingestion was performed.")
        elif args.country:
            codes={r[0] for r in __import__("app.catalog",fromlist=["COUNTRIES"]).COUNTRIES if r[0]==args.country.upper()}
            ids={r["source_id"] for r in catalogue() if r.get("country_code") in codes}
            candidates={k:v for k,v in candidates.items() if k in ids}
            if not candidates: raise SystemExit(f"No connector registered for country {args.country.upper()}; no ingestion was performed.")
        elif not args.all:
            raise SystemExit("Select --all, --country, or --source.")
        if not candidates:
            raise SystemExit("No runnable provider connectors are registered; no ingestion was performed.")
        from app.etl.orchestrator import run_many
        results=run_many(candidates,start_date=date.fromisoformat(args.start_date) if args.start_date else None,end_date=date.fromisoformat(args.end_date) if args.end_date else None)
        REPORTS.mkdir(exist_ok=True)
        run_report={"generated_at":datetime.now(timezone.utc).isoformat(),"results":results}
        (REPORTS/"latest_run.json").write_text(json.dumps(run_report,indent=2)+"\n")
        if any(row.get("status")=="Succeeded" for row in results):
            counts=export_database_samples()
            write_latest_run_report(run_report,counts)
        print(json.dumps(results,indent=2))

if __name__=="__main__": main()
