"""ASISA South African collective investment schemes fund-list connector."""
from __future__ import annotations

import hashlib
import json
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.config import timeout_seconds, user_agent
from app.connectors.base import BaseConnector

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / "south_africa"
REPORTS = ROOT / "reports"
URL = "https://www.asisa.org.za/media/jf5m4adc/list-of-cis-as-at-30-june-2026.xlsx"


def download_asisa_fund_list(*, refresh: bool = False) -> dict[str, Any]:
    RAW.mkdir(parents=True, exist_ok=True)
    target = RAW / "asisa_cis_fund_list_20260630.xlsx"
    downloaded = refresh or not target.exists()
    if downloaded:
        request = urllib.request.Request(URL, headers={"User-Agent": user_agent(), "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"})
        with urllib.request.urlopen(request, timeout=timeout_seconds()) as response:
            payload = response.read()
        if not payload.startswith(b"PK\x03\x04"):
            raise ValueError("ASISA response is not an XLSX file")
        target.write_bytes(payload)
    payload = target.read_bytes()
    manifest = {"source_id": "asisa_cis_fund_list", "url": URL,
        "path": str(target.relative_to(ROOT)), "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(), "downloaded": downloaded,
        "checked_at": datetime.now(timezone.utc).isoformat()}
    REPORTS.mkdir(parents=True, exist_ok=True)
    (REPORTS / "asisa_dataset_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


class AsisaConnector(BaseConnector):
    source_id = "asisa_cis_fund_list"
    country_code = "ZA"

    def __init__(self) -> None:
        self._records: list[dict[str, Any]] | None = None
        self._rejects: list[dict[str, Any]] = []

    def get_source_metadata(self) -> dict[str, Any]:
        return {"country_code": "ZA", "organization": "Association for Savings and Investment South Africa (ASISA)",
            "source_name": "CIS fund list as of 2026-06-30", "endpoint_url": URL,
            "license_status": "Unknown"}

    def _load(self) -> None:
        download_asisa_fund_list()
        workbook = load_workbook(RAW / "asisa_cis_fund_list_20260630.xlsx", read_only=True, data_only=True)
        sheet = workbook["CISFunds"]
        headers = [cell.value for cell in sheet[3]]
        if headers[:5] != ["Fund Code", "Fund Name", "Sector Code", "Sector Name", "Fund Type"]:
            raise ValueError(f"Unexpected ASISA fund list columns: {headers}")
        records = []
        self._rejects = []
        seen: set[str] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=4, max_col=9, values_only=True), start=4):
            code = str(row[0] or "").strip()
            name = str(row[1] or "").strip()
            if not code and not name:
                continue
            if not code or not name or code in seen:
                self._rejects.append({"row": row_number, "reason": "missing or duplicate fund code/name", "code": code})
                continue
            seen.add(code)
            records.append({"source_fund_id": code, "name": name,
                "domicile_country": "ZA", "currency": None, "isin": None, "_navs": []})
        workbook.close()
        self._records = records
        (REPORTS / "asisa_rejected_records.jsonl").write_text("".join(json.dumps(x, ensure_ascii=False) + "\n" for x in self._rejects))

    def check_availability(self) -> dict[str, Any]:
        self._load()
        return {"available": True, "fund_records": len(self._records or []), "rejected_records": len(self._rejects)}

    def extract_funds(self, start_date: date | None = None, end_date: date | None = None) -> Iterable[dict[str, Any]]:
        if self._records is None:
            self._load()
        return iter(self._records or [])

    def extract_nav_history(self, source_fund_id: str, start_date: date | None = None,
                            end_date: date | None = None) -> Iterable[dict[str, Any]]:
        return ()
