"""Scotia Canada official daily NAV workbook."""
from __future__ import annotations
import hashlib
from datetime import date
from pathlib import Path
from typing import Any, Iterable
from openpyxl import load_workbook
from app.connectors.base import BaseConnector

ROOT=Path(__file__).resolve().parents[2]
RAW=ROOT/'data/raw/canada/scotia_daily_nav.xlsx'
URL='https://2edd7e4a-a2d0-4023-bcb4-a5d649e18cd8-funds-sites.livedoc.transmissionmedia.ca/assets/static/daily-nav-exports/daily-nav-data-EN.xlsx'

def category(name: str) -> str:
    n=name.casefold()
    for key,label in [('money market','Money market'),('cash','Money market'),('bond','Fixed income'),('income','Income'),('balanced','Balanced'),('dividend','Equity income'),('equity','Equity'),('share','Equity'),('emerging markets','Emerging markets'),('international','International equity'),('global','Global equity'),('canadian','Canadian equity'),('us ','US equity'),('infrastructure','Real assets'),('real estate','Real estate'),('mortgage','Fixed income'),('index','Index')]:
        if key in n: return label
    return 'Other'

class ScotiaCanadaConnector(BaseConnector):
    source_id='scotia_canada_daily_nav'; country_code='CA'
    def __init__(self): self._records=None; self._navs={}
    def get_source_metadata(self): return {'country_code':'CA','organization':'Scotiabank / ScotiaFunds','source_name':'Daily Canadian mutual-fund NAV workbook','endpoint_url':URL,'license_status':'Public issuer data'}
    def _load(self):
        self._records=[]; self._navs={}
        wb=load_workbook(RAW,read_only=True,data_only=True); ws=wb.active
        for row in ws.iter_rows(min_row=2,values_only=True):
            name,series,currency,asof,nav,_,_,_,_,code=row
            name=str(name or '').strip(); series=str(series or '').strip(); currency=str(currency or '').strip().upper(); code=str(code or '').strip()
            if not name or not asof or not nav or currency not in {'CAD','USD'}: continue
            key=code+'|'+series+'|'+currency if code else hashlib.sha256(f'{name}|{series}|{currency}'.encode()).hexdigest()[:20]
            sid='CA-SCOTIA-'+key
            if not any(x['source_fund_id']==sid for x in self._records):
                is_etf=' ETF' in name.upper() or name.upper().endswith('ETF')
                self._records.append({'source_fund_id':sid,'name':f'{name} - Series {series}' if series else name,'legal_name':name,'domicile_country':'CA','currency':currency,'isin':None,'provider':'ScotiaFunds','manager':'Scotiabank','category':category(name),'fund_type':'ETF' if is_etf else 'Mutual fund','status':'Active','official_url':'https://www.scotiafunds.com/en/home/all-funds/mutual-fund','ticker':code or None})
            self._navs.setdefault(sid,[]).append({'source_fund_id':sid,'valuation_date':str(asof),'nav':str(nav),'currency':currency})
        wb.close()
    def check_availability(self): self._load(); return {'available':True,'fund_records':len(self._records),'nav_records':sum(map(len,self._navs.values()))}
    def extract_funds(self,start_date=None,end_date=None):
        if self._records is None:self._load()
        return iter(self._records)
    def extract_nav_history(self,source_fund_id,start_date=None,end_date=None):
        if self._records is None:self._load()
        return iter(r for r in self._navs.get(source_fund_id,[]) if (not start_date or date.fromisoformat(r['valuation_date'])>=start_date) and (not end_date or date.fromisoformat(r['valuation_date'])<=end_date))
