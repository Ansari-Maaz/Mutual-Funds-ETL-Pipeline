"""CSRC China public-fund product index connector."""
from __future__ import annotations

import hashlib
import json
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.config import timeout_seconds, user_agent
from app.connectors.base import BaseConnector

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw" / "china"
REPORTS = ROOT / "reports"
URL = "https://www.csrc.gov.cn/csrc/c101900/c1029655/1029655/files/%E5%85%AC%E5%8B%9F%E5%9F%BA%E9%87%91%E4%BA%A7%E5%93%81%E7%B4%A2%E5%BC%95%EF%BC%88%E6%88%AA%E8%87%B320260630%EF%BC%89.xlsx"


def download_csrc_index(*, refresh: bool = False) -> dict[str, Any]:
    RAW.mkdir(parents=True, exist_ok=True)
    target = RAW / "csrc_public_fund_index_20260630.xlsx"
    downloaded = refresh or not target.exists()
    if downloaded:
        request = urllib.request.Request(URL, headers={"User-Agent": user_agent(), "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"})
        with urllib.request.urlopen(request, timeout=timeout_seconds()) as response:
            payload = response.read()
        if not payload.startswith(b"PK\x03\x04"):
            raise ValueError("CSRC response is not an XLSX file")
        target.write_bytes(payload)
    payload = target.read_bytes()
    manifest = {"source_id": "csrc_public_fund_index", "url": URL,
        "path": str(target.relative_to(ROOT)), "bytes": len(payload),
        "sha256": hashlib.sha256(payload).hexdigest(), "downloaded": downloaded,
        "checked_at": datetime.now(timezone.utc).isoformat()}
    REPORTS.mkdir(parents=True, exist_ok=True)
    (REPORTS / "csrc_dataset_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return manifest


class CsrcConnector(BaseConnector):
    source_id = "csrc_public_fund_index"
    country_code = "CN"

    def __init__(self) -> None:
        self._records: list[dict[str, Any]] | None = None
        self._rejects: list[dict[str, Any]] = []

    def get_source_metadata(self) -> dict[str, Any]:
        return {"country_code": "CN", "organization": "China Securities Regulatory Commission (CSRC)",
            "source_name": "Public fund product index as of 2026-06-30", "endpoint_url": URL,
            "license_status": "Unknown"}

    def _load(self) -> None:
        download_csrc_index()
        workbook = load_workbook(RAW / "csrc_public_fund_index_20260630.xlsx", read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        header = [cell.value for cell in sheet[2]]
        if header[:4] != ["序号", "基金代码", "基金简称", "设立日期"]:
            raise ValueError(f"Unexpected CSRC product index columns: {header}")
        records = []
        self._rejects = []
        seen: set[str] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=3, max_col=4, values_only=True), start=3):
            _, raw_code, raw_name, _ = row
            if raw_code is None and raw_name is None:
                continue
            if raw_code is None or raw_name is None:
                self._rejects.append({"row": row_number, "reason": "missing fund code or name"})
                continue
            code = str(int(raw_code)).zfill(6) if isinstance(raw_code, (int, float)) else str(raw_code).strip()
            name = str(raw_name).strip()
            if not code or not name or code in seen:
                self._rejects.append({"row": row_number, "reason": "blank or duplicate code", "code": code})
                continue
            seen.add(code)
            records.append({"source_fund_id": code, "name": name,
                "domicile_country": "CN", "currency": None, "isin": None, "_navs": []})
        workbook.close()
        self._records = records
        (REPORTS / "csrc_rejected_records.jsonl").write_text("".join(json.dumps(x, ensure_ascii=False) + "\n" for x in self._rejects))

    def check_availability(self) -> dict[str, Any]:
        self._load()
        return {"available": True, "fund_records": len(self._records or []), "rejected_records": len(self._rejects)}

    def extract_funds(self, start_date: date | None = None, end_date: date | None = None) -> Iterable[dict[str, Any]]:
        if self._records is None:
            self._load()
        return iter(self._records or [])

    def extract_nav_history(self, source_fund_id: str, start_date: date | None = None,
                            end_date: date | None = None) -> Iterable[dict[str, Any]]:
        return ()
